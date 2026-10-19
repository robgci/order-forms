"""Create a blank intake workbook and a filled Handcraft example (requires openpyxl)."""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
CUSTOMER = json.loads((ROOT / 'data/handcraft.json').read_text())
OUT = ROOT / 'templates'
OUT.mkdir(exist_ok=True)

def sheet(wb, name, headers, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    ws.freeze_panes = 'A2'
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='29384D')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 36
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    return ws

def make(example):
    wb = Workbook()
    wb.remove(wb.active)
    guide = sheet(wb, 'Read me', ['Customer order form setup', 'Instructions'], [32, 110])
    notes = [
        ('One workbook per customer', 'Fill in Contacts and Products, save, then upload the workbook to the order-form conversation.'),
        ('Contacts', 'One row per person/location. Customer name, contact name/email, location and designated salesperson name/email are required.'),
        ('Dropdown', 'The form will list the contact people. Selecting a person automatically fills the location and uses that contact email for future customer delivery.'),
        ('Products', 'One row per item. Keep item numbers as text, exactly as supplied, including leading zeros and punctuation.'),
        ('Order unit and price', 'Choose pieces or dozen. Unit price must be the price for that order unit. Do not convert a per-piece price unless instructed.'),
        ('Pack sizes', 'Always enter case pack in PIECES and pallet pack in CASES PER PALLET. The three blue columns calculate the equivalents.'),
        ('Dozen example', 'Washcloth: 900 pieces/case and 8 cases/pallet = 75 dozen/case and 600 dozen/pallet. Other products stay in pieces.'),
        ('Full-case rounding', 'All positive quantities round UP to full cases. Zero stays zero. Fractional pallets are allowed. Dozen case packs must divide evenly by 12.'),
        ('Currency', 'Enter the three-letter currency for each customer (for example USD or CAD). Handcraft currency is intentionally blank because the source only showed dollar signs.'),
        ('Testing', 'Customer addresses are configuration data only. While testing, send only to rhorowitz@georgecourey.com. Customer/salesperson delivery must be explicitly enabled later.'),
        ('No secrets', 'Do not enter API keys, passwords or other credentials in this workbook.'),
        ('More rows', 'Add rows as needed. If adding Products rows, copy the three blue formulas down. Yellow cells are inputs; blue cells are calculated.'),
    ]
    for row in notes:
        guide.append(row)
        guide.row_dimensions[guide.max_row].height = 45
        for cell in guide[guide.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical='center')

    contacts = sheet(wb, 'Contacts', ['Customer name', 'Contact person', 'Customer contact email', 'Ship-to location', 'Salesperson', 'Salesperson email', 'Currency'], [30, 28, 48, 32, 25, 40, 14])
    products = sheet(wb, 'Products', ['Item number', 'Description', 'Order unit', 'Unit price', 'Case pack (pieces)', 'Pallet pack (cases)', 'Pallet pack (pieces)', 'Order units per case', 'Order units per pallet'], [22, 78, 16, 16, 23, 23, 23, 24, 25])
    for r in range(2, 22):
        for cell in contacts[r]:
            cell.fill = PatternFill('solid', fgColor='FFF8DB')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for cell in products[r]:
            cell.fill = PatternFill('solid', fgColor='EAF1FA' if cell.column >= 7 else 'FFF8DB')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        products.cell(r, 1).number_format = '@'
        products.cell(r, 4).number_format = '0.00'
        products.cell(r, 7, f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})')
        products.cell(r, 8, f'=IF(OR(C{r}="",E{r}=""),"",IF(C{r}="dozen",E{r}/12,E{r}))')
        products.cell(r, 9, f'=IF(OR(H{r}="",F{r}=""),"",H{r}*F{r})')
        products.row_dimensions[r].height = 32
        for col in range(7, 10): products.cell(r, col).number_format = '#,##0.###'
    unit = DataValidation(type='list', formula1='"pieces,dozen"')
    unit.errorTitle = 'Choose an order unit'; unit.error = 'Choose pieces or dozen.'; unit.showErrorMessage = True
    products.add_data_validation(unit); unit.add('C2:C1000')
    for col in ['E', 'F']:
        rule = DataValidation(type='whole', operator='between', formula1=1, formula2=1000000)
        rule.showErrorMessage = True; rule.error = 'Enter a positive whole number.'
        products.add_data_validation(rule); rule.add(f'{col}2:{col}1000')
    price = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0)
    price.showErrorMessage = True; price.error = 'Enter a nonnegative price per order unit.'
    products.add_data_validation(price); price.add('D2:D1000')
    products['E1'].comment = Comment('Always pieces, including products ordered in dozens.', 'Order forms')
    products['F1'].comment = Comment('Number of complete cases per pallet.', 'Order forms')
    products['D1'].comment = Comment('Price per piece or per dozen, matching Order unit.', 'Order forms')
    for ws, table_name, end in [(contacts, 'CustomerContacts', 'G'), (products, 'CustomerProducts', 'I')]:
        table = Table(displayName=table_name, ref=f'A1:{end}21')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False)
        ws.add_table(table)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.print_title_rows = '1:1'
    if example:
        for r, contact in enumerate(CUSTOMER['contacts'], 2):
            values = [CUSTOMER['name'], contact['name'], contact['email'], contact['location'], CUSTOMER['salesRep'], CUSTOMER['salesEmail']]
            for col, value in enumerate(values, 1): contacts.cell(r, col, value)
        for r, product in enumerate(CUSTOMER['products'], 2):
            values = [product['id'], product['name'], 'dozen' if product['piecesPerUnit'] == 12 else 'pieces', product['price'], product['piecesPerCase'], product['casesPerPallet']]
            for col, value in enumerate(values, 1): products.cell(r, col, value)
            products.row_dimensions[r].height = 64
    wb.active = 1
    path = OUT / ('Handcraft-Setup-Example.xlsx' if example else 'Customer-Order-Form-Template.xlsx')
    wb.save(path)
    check = load_workbook(path)
    assert check.sheetnames == ['Read me', 'Contacts', 'Products']
    assert check['Products']['H2'].data_type == 'f'
    if example:
        assert check['Contacts']['B4'].value == 'Clinton Wright'
        assert check['Products']['E5'].value == 900
        assert check['Products']['C5'].value == 'dozen'
    print(path)

make(False)
make(True)
